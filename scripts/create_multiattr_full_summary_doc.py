from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "reports" / "多属性动态范围过滤DSG完整实验汇总.docx"

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
MUTED = RGBColor(90, 90, 90)
HEADER_FILL = "E8EEF5"
LIGHT_FILL = "F4F6F9"
WIDTH = 9360


def set_run(run, size=10.5, bold=False, color=None, font="Microsoft YaHei"):
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), font)
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    node = tc_pr.find(qn("w:shd"))
    if node is None:
        node = OxmlElement("w:shd")
        tc_pr.append(node)
    node.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=70, start=100, bottom=70, end=100):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def geometry(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, widths, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    geometry(table, widths)
    for idx, header in enumerate(headers):
        table.cell(0, idx).text = header
        shade(table.cell(0, idx), HEADER_FILL)
        for run in table.cell(0, idx).paragraphs[0].runs:
            set_run(run, size=9.2, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = str(value)
            for run in cells[idx].paragraphs[0].runs:
                set_run(run, size=9.0)
    return table


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.12
    set_run(p.add_run(text))


def number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.12
    set_run(p.add_run(text))


def code(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.space_after = Pt(1)
    set_run(p.add_run(text), size=8.8, font="Consolas")


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.7)
section.bottom_margin = Inches(0.7)
section.left_margin = Inches(0.8)
section.right_margin = Inches(0.8)
section.header_distance = Inches(0.35)
section.footer_distance = Inches(0.35)

normal = doc.styles["Normal"]
normal.font.name = "Microsoft YaHei"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
normal.font.size = Pt(10)
normal.paragraph_format.space_after = Pt(4)
normal.paragraph_format.line_spacing = 1.12

for name, size, color, before, after in (
    ("Heading 1", 15, BLUE, 12, 6),
    ("Heading 2", 12.5, BLUE, 8, 4),
    ("Heading 3", 11, DARK_BLUE, 6, 3),
):
    style = doc.styles[name]
    style.font.name = "Microsoft YaHei"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = color
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

title = doc.add_paragraph()
title.paragraph_format.space_after = Pt(3)
set_run(title.add_run("多属性动态范围过滤 DSG 完整实验汇总"), size=21, bold=True, color=DARK_BLUE)

subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(10)
set_run(
    subtitle.add_run("论文代码扩展、实现过程、实验结果与当前结论 | 2026-06-10"),
    size=9.5,
    color=MUTED,
)

lead = doc.add_table(rows=1, cols=1)
geometry(lead, [WIDTH])
shade(lead.cell(0, 0), LIGHT_FILL)
p = lead.cell(0, 0).paragraphs[0]
p.paragraph_format.space_after = Pt(0)
set_run(
    p.add_run(
        "最终推荐：使用主属性 DSG 进行导航，非主属性仅过滤最终候选。"
        "推荐参数 M=16、ef_construction=100、ef_max=200、search_ef=16。"
        "10K 实验 Recall@10=1.0，QPS=59,406。"
    ),
    bold=True,
)

doc.add_heading("1. 原始论文与代码约束", level=1)
bullet(doc, "论文问题定义为单个有全序的数值属性；查询是一个一维范围。")
bullet(doc, "论文“矩形树”的二维坐标是查询左端点和右端点，不是两个属性。")
bullet(doc, "原代码默认 label 同时承担点 ID、属性值和属性 rank，查询接口仅支持 pair<int,int>。")
bullet(doc, "原边标签由 left_lower/left_upper/right_lower/right_upper 表示单属性 envelope。")

doc.add_heading("2. 总体扩展思路", level=1)
number(doc, "保留原 DSG 的单属性压缩结构，选择 attr0 作为主属性。")
number(doc, "将多属性过滤条件封装为 MultiRangeQuery。")
number(doc, "主属性范围转换为 rank 区间，沿用 DSG envelope 导航。")
number(doc, "图搜索允许不满足非主属性的节点作为桥接节点。")
number(doc, "仅将满足全部属性范围的访问节点加入结果候选，再按向量距离重排。")
number(doc, "使用 exact 多属性扫描作为真值，持续对比 Recall、QPS、hops 和距离计算数。")

doc.add_heading("3. 代码新增与改动", level=1)
code_rows = [
    ("include/utils/filter_query.h", "新增", "定义 RangeBound、MultiRangeQuery。"),
    ("include/base_index.h", "修改", "新增多属性 rangeSearch 虚接口。"),
    ("include/utils/data_wrapper.h", "修改", "新增 attrs、attr_rank、rank_to_label、primary_attr 及过滤接口。"),
    ("src/utils/data_wrapper.cc", "修改", "实现 CSV 属性读取、属性排序、值域转 rank、passFilter。"),
    ("include/dsg.h", "修改", "声明多属性图查询和 rangeSearchMultiExact。"),
    ("src/dsg.cc", "修改", "实现 exact 查询、专用多属性 DSG 搜索、自适应/禁用 MBR 消融。"),
    ("apps/static/query_multi_filter_index.cc", "新增", "逐条打印 expected、exact、graph，用于 smoke test。"),
    ("apps/static/query_multi_filter_benchmark.cc", "新增", "按 profile 汇总 Recall@K、QPS、hops、distance。"),
    ("apps/CMakeLists.txt", "修改", "注册两个多属性测试可执行目标。"),
    ("scripts/generate_multiattr_smoke_data.py", "新增", "生成 N=32、dim=4、3 属性、4 查询的小数据。"),
    ("scripts/generate_multiattr_benchmark_data.py", "新增", "生成 N=10K、dim=32、3 属性、100 查询及 exact Top-10。"),
]
add_table(doc, ("文件", "类型", "功能"), [3000, 1000, 5360], code_rows)

doc.add_heading("4. 查询实现演进", level=1)
evolution_rows = [
    ("Exact smoke", "全量扫描 + passFilter", "验证数据读取和多属性过滤正确性。"),
    ("Graph + fallback", "DSG 候选不足时补扫主属性范围", "32 点全部对齐 exact，但包含扫描兜底。"),
    ("Dedicated graph", "搜索过程直接收集访问节点", "去除 fallback；ef>=16 时 32 点 Recall@3=1。"),
    ("Strict loose MBR", "非主属性不匹配则拒绝边", "距离计算下降，但宽查询断连。"),
    ("Adaptive MBR", "选择率窄时启用硬 MBR", "32 点可用；10K Recall@10 仅 0.95。"),
    ("Bridge edges", "每节点允许 2 条不匹配边", "增加计算但 Recall 不变。"),
    ("No MBR", "非主属性仅用于结果过滤", "10K Recall@10=1.0，性能最佳。"),
]
add_table(doc, ("阶段", "策略", "效果"), [1800, 3000, 4560], evolution_rows)

doc.add_heading("5. 测试数据", level=1)
data_rows = [
    ("multiattr_smoke", "32", "4", "4", "3", "功能正确性、逐条结果对照"),
    ("multiattr_10k", "10,000", "32", "100", "10", "批量 Recall/QPS 与消融实验"),
]
add_table(doc, ("数据集", "N", "维度", "查询数", "Top-K", "用途"), [1700, 900, 900, 1100, 900, 3860], data_rows)
bullet(doc, "attr0 == label，保证与当前主属性 rank 假设兼容。")
bullet(doc, "attr1、attr2 使用独立随机排列，避免属性相关性导致虚假收益。")
bullet(doc, "10K 查询分为 narrow、medium、mixed、broad 四类，每类 25 条。")

doc.add_heading("6. 关键实验结果", level=1)
doc.add_heading("6.1 32 点功能验证", level=2)
bullet(doc, "Exact 查询 4/4 与预生成 expected Top-3 完全一致。")
bullet(doc, "专用图搜索在 ef>=16 时 4/4 与 exact 完全一致。")
bullet(doc, "严格 loose-MBR 下 q3 固定错误，证明硬过滤造成子图断连。")

doc.add_heading("6.2 10K 参数实验", level=2)
ef_rows = [
    ("16", "0.9500", "55,162", "18.16", "73.61"),
    ("32", "0.9500", "32,004", "31.59", "99.52"),
    ("64", "0.9500", "22,099", "54.45", "134.78"),
    ("128", "0.9500", "13,425", "100.35", "187.17"),
    ("256", "0.9500", "8,555", "168.64", "262.43"),
]
add_table(doc, ("search_ef", "Recall@10", "QPS", "Avg hops", "Avg dist"), [1600, 1800, 1800, 2000, 2160], ef_rows)
bullet(doc, "Adaptive MBR 的 Recall 不随 ef 增大，说明问题不是搜索宽度，而是硬边过滤。")

doc.add_heading("6.3 索引质量与桥接消融", level=2)
bullet(doc, "M 从 16 增加到 32，Recall 仍为 0.95；索引度数不是瓶颈。")
bullet(doc, "每节点增加 2 条桥接边，Recall 不变，且距离计算增加。")
bullet(doc, "完全关闭非主属性硬边过滤后，M16/M32 均达到 Recall@10=1.0。")

doc.add_heading("6.4 最终配置对比", level=2)
final_rows = [
    ("Adaptive MBR, M16, ef=16", "0.9500", "55,162", "73.61"),
    ("No MBR, M16, ef=16", "1.0000", "59,406", "89.80"),
    ("No MBR, M32, ef=16", "1.0000", "58,885", "91.09"),
]
add_table(doc, ("配置", "Recall@10", "QPS", "Avg dist"), [3900, 1800, 1600, 2060], final_rows)

doc.add_heading("7. 当前结论", level=1)
bullet(doc, "第一阶段可行方案是“主属性导航 + 多属性结果过滤”。")
bullet(doc, "非主属性 loose-MBR 不应作为硬边过滤，否则会损害图连通性。")
bullet(doc, "在线 MBR rank 判断与分支存在额外 CPU 开销，本实验中 QPS 反而低于 no-MBR。")
bullet(doc, "M16 已达到完整召回并优于 M32，继续增大索引度数没有收益。")
bullet(doc, "当前推荐基线：M16、ef_construction=100、ef_max=200、search_ef=16、no-MBR traversal。")

doc.add_heading("8. 备份与当前代码状态", level=1)
backup_rows = [
    ("src/dsg.cc.graph_fallback_ok", "图搜索 + 主属性范围 fallback 正确性版本"),
    ("src/dsg.cc.strict_loose_mbr", "严格 loose-MBR 消融版本"),
    ("src/dsg.cc.adaptive_mbr_ok", "32 点自适应 MBR 版本"),
    ("src/dsg.cc.adaptive_mbr_10k_ok", "10K 自适应 MBR 基线"),
    ("src/dsg.cc.before_no_mbr_test", "关闭 MBR 前的回滚点"),
    ("src/dsg.cc.no_mbr_recall100", "10K Recall@10=1.0 版本"),
    ("src/dsg.cc.multiattr_baseline_ok", "当前推荐 DSG 实现备份"),
    ("include/dsg.h.multiattr_baseline_ok", "当前推荐接口备份"),
    ("query_multi_filter_benchmark.cc.multiattr_baseline_ok", "当前 benchmark 备份"),
]
add_table(doc, ("备份文件", "用途"), [4100, 5260], backup_rows)

doc.add_heading("8.1 建议立即执行的备份命令", level=2)
code(doc, "cp src/dsg.cc src/dsg.cc.multiattr_baseline_ok")
code(doc, "cp include/dsg.h include/dsg.h.multiattr_baseline_ok")
code(
    doc,
    "cp apps/static/query_multi_filter_benchmark.cc "
    "apps/static/query_multi_filter_benchmark.cc.multiattr_baseline_ok",
)
code(doc, "cp -r data/multiattr_10k data/multiattr_10k.baseline")

doc.add_heading("9. 下一阶段", level=1)
bullet(doc, "实现 soft MBR：匹配边优先展开，不匹配边保留为低优先级桥接路径。")
bullet(doc, "将 MBR 选择率和 bridge penalty 作为查询参数，做 Recall/QPS 曲线。")
bullet(doc, "验证动态插入后的多属性查询、recompress 和插入耗时。")
bullet(doc, "使用真实多属性数据验证，并补充索引大小与内存统计。")
bullet(doc, "完成多属性元数据与索引 save/load 持久化。")

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
set_run(footer.add_run("Multi-Attribute DSG Experiment Summary"), size=8, color=MUTED)

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
