from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "reports" / "多属性动态范围过滤DSG阶段总结.docx"

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
MUTED = RGBColor(90, 90, 90)
HEADER_FILL = "E8EEF5"
LIGHT_FILL = "F4F6F9"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_fixed_table_geometry(table, widths_dxa):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_run(run, size=11, bold=False, color=None, font="Microsoft YaHei"):
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), font)
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def add_bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.15
    set_run(paragraph.add_run(text))
    return paragraph


def add_number(doc, text):
    paragraph = doc.add_paragraph(style="List Number")
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.15
    set_run(paragraph.add_run(text))
    return paragraph


def add_code_line(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.2)
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(text)
    set_run(run, size=9, font="Consolas")
    return paragraph


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.75)
section.bottom_margin = Inches(0.75)
section.left_margin = Inches(0.85)
section.right_margin = Inches(0.85)
section.header_distance = Inches(0.35)
section.footer_distance = Inches(0.35)

normal = doc.styles["Normal"]
normal.font.name = "Microsoft YaHei"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
normal.font.size = Pt(10.5)
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.line_spacing = 1.15

for name, size, color, before, after in (
    ("Heading 1", 16, BLUE, 14, 7),
    ("Heading 2", 13, BLUE, 10, 5),
    ("Heading 3", 11.5, DARK_BLUE, 7, 3),
):
    style = doc.styles[name]
    style.font.name = "Microsoft YaHei"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = color
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

title = doc.add_paragraph()
title.paragraph_format.space_after = Pt(4)
run = title.add_run("多属性动态范围过滤 DSG 阶段总结")
set_run(run, size=22, bold=True, color=DARK_BLUE)

subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(12)
set_run(subtitle.add_run("主属性导航 + 多属性结果过滤 | 2026-06-10"), size=10, color=MUTED)

lead = doc.add_table(rows=1, cols=1)
set_fixed_table_geometry(lead, [9360])
set_cell_shading(lead.cell(0, 0), LIGHT_FILL)
p = lead.cell(0, 0).paragraphs[0]
p.paragraph_format.space_after = Pt(0)
set_run(
    p.add_run(
        "当前推荐基线：M=16，ef_construction=100，ef_max=200，search_ef=16。"
        "主属性 DSG 负责图导航，非主属性仅用于结果过滤；不使用非主属性硬边过滤。"
    ),
    bold=True,
)

doc.add_heading("1. 已完成工作", level=1)
add_bullet(doc, "完成论文与代码对照：论文矩形树是单属性查询左右端点空间，不是多属性矩形。")
add_bullet(doc, "新增 MultiRangeQuery、属性读取、属性 rank、passFilter 和多属性精确查询。")
add_bullet(doc, "实现多属性 DSG 查询：主属性 envelope 约束图遍历，全部属性约束最终结果。")
add_bullet(doc, "实现 32 点 smoke test，exact、graph+fallback、纯图搜索均完成验证。")
add_bullet(doc, "生成 10K/32维/3属性/100查询中型数据，并实现批量 Recall/QPS benchmark。")

doc.add_heading("2. 关键文件", level=1)
table = doc.add_table(rows=1, cols=2)
table.style = "Table Grid"
set_fixed_table_geometry(table, [3000, 6360])
headers = table.rows[0].cells
headers[0].text = "文件"
headers[1].text = "作用"
for cell in headers:
    set_cell_shading(cell, HEADER_FILL)
    for run in cell.paragraphs[0].runs:
        set_run(run, bold=True)

files = [
    ("include/utils/filter_query.h", "RangeBound、MultiRangeQuery 定义"),
    ("include/utils/data_wrapper.h", "属性表、rank、过滤接口声明"),
    ("src/utils/data_wrapper.cc", "属性读取、rank 构建、值域转 rank、passFilter"),
    ("include/dsg.h / src/dsg.cc", "多属性 exact 与 DSG 图搜索实现"),
    ("apps/static/query_multi_filter_index.cc", "小数据逐条正确性测试"),
    ("apps/static/query_multi_filter_benchmark.cc", "中型数据批量指标汇总"),
    ("scripts/generate_multiattr_smoke_data.py", "32 点 smoke 数据生成"),
    ("scripts/generate_multiattr_benchmark_data.py", "10K 中型数据及精确真值生成"),
]
for path, purpose in files:
    cells = table.add_row().cells
    cells[0].text = path
    cells[1].text = purpose
    for cell in cells:
        for run in cell.paragraphs[0].runs:
            set_run(run, size=9.5)

doc.add_heading("3. 实验结论", level=1)
result_table = doc.add_table(rows=1, cols=4)
result_table.style = "Table Grid"
set_fixed_table_geometry(result_table, [3000, 1800, 2100, 2460])
for idx, text in enumerate(("配置", "Recall@10", "QPS", "平均距离计算")):
    result_table.cell(0, idx).text = text
    set_cell_shading(result_table.cell(0, idx), HEADER_FILL)
    for run in result_table.cell(0, idx).paragraphs[0].runs:
        set_run(run, bold=True, size=9.5)

results = [
    ("Adaptive MBR, M16, ef=16", "0.9500", "55,162", "73.61"),
    ("No MBR, M16, ef=16", "1.0000", "59,406", "89.80"),
    ("No MBR, M32, ef=16", "1.0000", "58,885", "91.09"),
]
for row in results:
    cells = result_table.add_row().cells
    for idx, text in enumerate(row):
        cells[idx].text = text
        for run in cells[idx].paragraphs[0].runs:
            set_run(run, size=9.5)

add_bullet(doc, "M32 未提升召回或吞吐，M16 足够。")
add_bullet(doc, "硬 loose-MBR 虽减少距离计算，但会切断导航路径，Recall 降至 0.95。")
add_bullet(doc, "关闭非主属性边过滤后，四类查询均达到 Recall@10=1.0。")
add_bullet(doc, "search_ef=16 已达到最佳 Recall/QPS 平衡；继续增大 ef 只增加开销。")

doc.add_heading("4. 当前算法", level=1)
add_number(doc, "将主属性值范围转换为主属性 rank 区间。")
add_number(doc, "只使用主属性 DSG envelope 判断边是否可用于导航。")
add_number(doc, "搜索过程中收集访问节点；不因非主属性不匹配而拒绝桥接节点。")
add_number(doc, "仅对满足全部属性范围的节点加入有效候选集。")
add_number(doc, "按真实向量距离重排并返回 Top-K。")

doc.add_heading("5. 备份与回滚", level=1)
backup_table = doc.add_table(rows=1, cols=2)
backup_table.style = "Table Grid"
set_fixed_table_geometry(backup_table, [4200, 5160])
for idx, text in enumerate(("建议备份文件", "用途")):
    backup_table.cell(0, idx).text = text
    set_cell_shading(backup_table.cell(0, idx), HEADER_FILL)
    for run in backup_table.cell(0, idx).paragraphs[0].runs:
        set_run(run, bold=True, size=9.5)

backups = [
    ("src/dsg.cc.graph_fallback_ok", "图搜索不足时主属性范围补扫的正确性版本"),
    ("src/dsg.cc.adaptive_mbr_ok", "32 点自适应 MBR 验证版本"),
    ("src/dsg.cc.strict_loose_mbr", "严格 loose-MBR 消融版本"),
    ("src/dsg.cc.adaptive_mbr_10k_ok", "10K 自适应 MBR 基线"),
    ("src/dsg.cc.no_mbr_recall100", "10K Recall=1.0 的 no-MBR 版本"),
    ("src/dsg.cc.multiattr_baseline_ok", "当前推荐多属性基线"),
    ("include/dsg.h.multiattr_baseline_ok", "当前推荐头文件备份"),
    ("apps/static/query_multi_filter_benchmark.cc.multiattr_baseline_ok", "批量 benchmark 备份"),
]
for name, purpose in backups:
    cells = backup_table.add_row().cells
    cells[0].text = name
    cells[1].text = purpose
    for cell in cells:
        for run in cell.paragraphs[0].runs:
            set_run(run, size=9.3)

doc.add_paragraph()
add_code_line(doc, "cp src/dsg.cc src/dsg.cc.multiattr_baseline_ok")
add_code_line(doc, "cp include/dsg.h include/dsg.h.multiattr_baseline_ok")
add_code_line(
    doc,
    "cp apps/static/query_multi_filter_benchmark.cc "
    "apps/static/query_multi_filter_benchmark.cc.multiattr_baseline_ok",
)

doc.add_heading("6. 下一步", level=1)
add_bullet(doc, "以 No-MBR M16/ef=16 作为正式 baseline。")
add_bullet(doc, "研究 soft MBR：匹配边优先，不匹配边保留为低优先级桥接路径。")
add_bullet(doc, "在真实多属性数据上验证 Recall、QPS、索引大小和动态插入开销。")
add_bullet(doc, "完成后再扩展 save/load，使多属性元数据与索引一并持久化。")

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
set_run(footer.add_run("Dynamic Range-Filtering ANNS - Stage Summary"), size=8.5, color=MUTED)

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
